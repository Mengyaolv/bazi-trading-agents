"""Report export helpers for Bazi analysis results."""

import io
import json
import os
import re
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple

from utils.time_utils import current_analysis_date, current_analysis_datetime
from utils.visualization import build_visualization_data
from utils.web_support import extract_intermediate_reports, sanitize_error


REPORT_TITLE = "\u516b\u5b57\u7efc\u5408\u5206\u6790\u62a5\u544a"
ETHICAL_STATEMENT = (
    "\u4f26\u7406\u58f0\u660e\uff1a\u672c\u62a5\u544a\u4ec5\u7528\u4e8e\u6587\u5316\u7814\u7a76\u3001"
    "\u81ea\u6211\u53cd\u601d\u548c\u51b3\u7b56\u53c2\u8003\uff0c\u4e0d\u6784\u6210\u533b\u7597\u3001"
    "\u6cd5\u5f8b\u3001\u6295\u8d44\u6216\u5176\u4ed6\u4e13\u4e1a\u610f\u89c1\u3002"
    "\u8bf7\u7ed3\u5408\u73b0\u5b9e\u8bc1\u636e\u548c\u4e13\u4e1a\u5efa\u8bae\u8c28\u614e\u4f7f\u7528\u3002"
)
REQUIRED_SECTIONS = [
    "\u547d\u76d8\u6458\u8981",
    "\u4f20\u7edf\u547d\u7406\u5206\u6790",
    "\u73b0\u4ee3\u5fc3\u7406\u5b66\u8f6c\u8bd1",
    "\u5f53\u524d\u9636\u6bb5\u7b56\u7565",
    "\u884c\u52a8\u5efa\u8bae",
    "\u98ce\u9669\u4e0e\u8fb9\u754c",
]
FORMAT_EXTENSIONS = {
    "markdown": "md",
    "md": "md",
    "json": "json",
    "word": "docx",
    "docx": "docx",
    "pdf": "pdf",
}


class ExportDependencyError(RuntimeError):
    """Raised when an optional export dependency is not available."""


def normalize_final_state(payload: Dict[str, Any]) -> Dict[str, Any]:
    """Normalize graph logs, batch records, and raw final_state dictionaries."""
    if not isinstance(payload, dict):
        return {}

    if isinstance(payload.get("final_state"), dict):
        state = dict(payload["final_state"])
        for key in (
            "id",
            "name",
            "birth_info",
            "analysis_date",
            "analysis_datetime",
            "final_analysis",
            "quality_report",
            "visualization_data",
            "prompt_version",
            "llm_provider",
            "llm_model",
            "metadata",
        ):
            if key not in state and key in payload:
                state[key] = payload[key]
    else:
        state = dict(payload)

    if not state.get("final_analysis") and payload.get("final_analysis"):
        state["final_analysis"] = payload["final_analysis"]
    if not state.get("visualization_data"):
        state["visualization_data"] = build_visualization_data(state)
    return state


def normalize_export_format(export_format: str) -> Tuple[str, str]:
    """Return canonical format name and file extension."""
    key = str(export_format or "").strip().lower()
    if key not in FORMAT_EXTENSIONS:
        raise ValueError(f"Unsupported export format: {export_format}")
    canonical = "markdown" if key == "md" else "word" if key == "docx" else key
    return canonical, FORMAT_EXTENSIONS[key]


def normalize_export_formats(formats: Any) -> List[str]:
    """Normalize a comma-separated string or iterable of export formats."""
    if isinstance(formats, str):
        raw_items = [item.strip() for item in formats.split(",")]
    elif formats is None:
        raw_items = ["markdown", "json"]
    else:
        raw_items = [str(item).strip() for item in formats]
    normalized = []
    for item in raw_items:
        if not item:
            continue
        canonical, _ext = normalize_export_format(item)
        if canonical not in normalized:
            normalized.append(canonical)
    return normalized or ["markdown", "json"]


def safe_filename(value: Any, fallback: str = "report", max_length: int = 80) -> str:
    """Return a filesystem-safe filename fragment."""
    text = str(value or "").strip()
    text = re.sub(r"[^A-Za-z0-9\u4e00-\u9fff_-]+", "_", text).strip("_")
    return (text[:max_length] or fallback).strip("_") or fallback


def build_export_filename(
    final_state: Dict[str, Any],
    case_id: str = "",
    ext: str = "md",
    mode: str = "",
) -> str:
    """Build the default report filename."""
    state = normalize_final_state(final_state)
    identifier = case_id or state.get("id") or state.get("birth_info") or "case"
    analysis_date = state.get("analysis_date") or current_analysis_date()
    mode_suffix = f"_{safe_filename(mode)}" if mode else ""
    return f"bazi_analysis_{safe_filename(identifier)}_{safe_filename(analysis_date)}{mode_suffix}.{ext}"


def build_export_payload(
    final_state: Dict[str, Any],
    case_id: str = "",
    mode: str = "simple",
) -> Dict[str, Any]:
    """Build a structured payload shared by Markdown, Word, PDF, and JSON."""
    state = normalize_final_state(final_state)
    mode = _normalize_mode(mode)
    visualization_data = state.get("visualization_data") or build_visualization_data(state)
    return {
        "title": REPORT_TITLE,
        "generated_at": current_analysis_datetime(),
        "case_id": case_id or state.get("id", ""),
        "birth_info": state.get("birth_info", ""),
        "analysis_date": state.get("analysis_date", ""),
        "analysis_datetime": state.get("analysis_datetime", ""),
        "mode": mode,
        "ethical_statement": ETHICAL_STATEMENT,
        "sections": _build_required_sections(state),
        "final_analysis": state.get("final_analysis", ""),
        "visualization_data": visualization_data,
        "intermediate_reports": extract_intermediate_reports(state),
        "quality_report": state.get("quality_report") or {},
        "raw_state": state,
    }


def render_markdown_report(
    final_state: Dict[str, Any],
    case_id: str = "",
    mode: str = "simple",
) -> str:
    """Render a report as Markdown."""
    payload = build_export_payload(final_state, case_id=case_id, mode=mode)
    professional = payload["mode"] == "professional"
    lines = [
        f"# {payload['title']}",
        "",
        f"- \u751f\u6210\u65f6\u95f4: {payload['generated_at']}",
        f"- \u5206\u6790\u65e5\u671f: {payload.get('analysis_date') or 'N/A'}",
    ]
    if payload.get("analysis_datetime"):
        lines.append(f"- \u5f53\u524d\u5206\u6790\u65f6\u95f4: {payload['analysis_datetime']}")
    if payload.get("case_id"):
        lines.append(f"- \u6848\u4f8b\u7f16\u53f7: {payload['case_id']}")
    if payload.get("birth_info"):
        lines.append(f"- \u51fa\u751f\u4fe1\u606f: {payload['birth_info']}")

    lines.extend(["", "## \u4f26\u7406\u58f0\u660e", "", payload["ethical_statement"], ""])

    for section in payload["sections"]:
        lines.extend([f"## {section['title']}", "", section["content"] or "\u6682\u65e0\u5185\u5bb9\u3002", ""])

    if payload.get("final_analysis"):
        lines.extend(["## \u6700\u7ec8\u7efc\u5408\u62a5\u544a", "", payload["final_analysis"], ""])

    lines.extend(_render_visualization_markdown(payload.get("visualization_data") or {}))

    if professional:
        lines.extend(_render_intermediate_markdown(payload.get("intermediate_reports") or []))
        lines.extend(_render_quality_markdown(payload.get("quality_report") or {}))
        lines.extend(
            [
                "## \u539f\u59cb JSON \u9644\u5f55",
                "",
                "```json",
                json.dumps(payload.get("raw_state") or {}, ensure_ascii=False, indent=2, default=str),
                "```",
                "",
            ]
        )

    return "\n".join(lines).rstrip() + "\n"


def export_report_bytes(
    final_state: Dict[str, Any],
    export_format: str,
    case_id: str = "",
    mode: str = "simple",
    font_path: Optional[str] = None,
) -> bytes:
    """Render one export format to bytes."""
    canonical, _ext = normalize_export_format(export_format)
    if canonical == "markdown":
        return render_markdown_report(final_state, case_id=case_id, mode=mode).encode("utf-8")
    if canonical == "json":
        payload = build_export_payload(final_state, case_id=case_id, mode=mode)
        return json.dumps(payload, ensure_ascii=False, indent=2, default=str).encode("utf-8")
    if canonical == "word":
        markdown = render_markdown_report(final_state, case_id=case_id, mode=mode)
        return _render_docx_bytes(markdown)
    if canonical == "pdf":
        markdown = render_markdown_report(final_state, case_id=case_id, mode=mode)
        return _render_pdf_bytes(markdown, font_path=font_path)
    raise ValueError(f"Unsupported export format: {export_format}")


def export_report(
    final_state: Dict[str, Any],
    output_dir: str,
    formats: Any = ("markdown", "json"),
    case_id: str = "",
    mode: str = "simple",
    font_path: Optional[str] = None,
) -> Dict[str, Dict[str, str]]:
    """Export a report to files, returning paths and non-fatal errors."""
    output_root = Path(output_dir)
    output_root.mkdir(parents=True, exist_ok=True)

    result: Dict[str, Dict[str, str]] = {"paths": {}, "errors": {}}
    state = normalize_final_state(final_state)
    for export_format in normalize_export_formats(formats):
        try:
            canonical, ext = normalize_export_format(export_format)
            file_name = build_export_filename(state, case_id=case_id, ext=ext, mode=mode)
            path = _unique_path(output_root / file_name)
            path.write_bytes(
                export_report_bytes(
                    state,
                    canonical,
                    case_id=case_id,
                    mode=mode,
                    font_path=font_path,
                )
            )
            result["paths"][canonical] = str(path)
        except Exception as exc:
            result["errors"][str(export_format)] = sanitize_error(exc)
    return result


def export_markdown(final_state: Dict[str, Any], output_path: str, **kwargs: Any) -> str:
    """Export Markdown to a specific path."""
    Path(output_path).write_bytes(export_report_bytes(final_state, "markdown", **kwargs))
    return output_path


def export_json(final_state: Dict[str, Any], output_path: str, **kwargs: Any) -> str:
    """Export JSON to a specific path."""
    Path(output_path).write_bytes(export_report_bytes(final_state, "json", **kwargs))
    return output_path


def export_docx(final_state: Dict[str, Any], output_path: str, **kwargs: Any) -> str:
    """Export Word to a specific path."""
    Path(output_path).write_bytes(export_report_bytes(final_state, "word", **kwargs))
    return output_path


def export_pdf(final_state: Dict[str, Any], output_path: str, **kwargs: Any) -> str:
    """Export PDF to a specific path."""
    Path(output_path).write_bytes(export_report_bytes(final_state, "pdf", **kwargs))
    return output_path


def _normalize_mode(mode: str) -> str:
    value = str(mode or "simple").strip().lower()
    if value not in {"simple", "professional"}:
        raise ValueError("Export mode must be 'simple' or 'professional'")
    return value


def _build_required_sections(state: Dict[str, Any]) -> List[Dict[str, str]]:
    debate_state = state.get("bazi_debate_state") or {}
    sections = {
        REQUIRED_SECTIONS[0]: state.get("chart_report") or state.get("final_analysis", ""),
        REQUIRED_SECTIONS[1]: debate_state.get("traditional_history")
        or "\n\n".join(
            item
            for item in (
                state.get("day_master_report", ""),
                state.get("ten_gods_report", ""),
                state.get("pattern_report", ""),
                state.get("useful_god_report", ""),
            )
            if item
        )
        or state.get("final_analysis", ""),
        REQUIRED_SECTIONS[2]: state.get("psychology_translation")
        or debate_state.get("psychology_history")
        or state.get("final_analysis", ""),
        REQUIRED_SECTIONS[3]: state.get("luck_cycle_report") or state.get("final_analysis", ""),
        REQUIRED_SECTIONS[4]: state.get("final_analysis", ""),
        REQUIRED_SECTIONS[5]: ETHICAL_STATEMENT,
    }
    return [{"title": title, "content": str(sections.get(title, "") or "")} for title in REQUIRED_SECTIONS]


def _render_visualization_markdown(data: Dict[str, Any]) -> List[str]:
    lines = ["## \u6838\u5fc3\u56fe\u8868", ""]
    table_specs = [
        ("\u56db\u67f1\u8868", data.get("pillars") or []),
        ("\u4e94\u884c\u5206\u5e03", data.get("element_distribution") or []),
        ("\u5341\u795e\u5206\u5e03", data.get("ten_gods_distribution") or []),
        ("\u5927\u8fd0\u65f6\u95f4\u8f74", data.get("luck_cycles") or []),
    ]
    rendered_any = False
    for title, rows in table_specs:
        if not rows:
            continue
        rendered_any = True
        lines.extend([f"### {title}", "", markdown_table(rows), ""])

    if not rendered_any:
        lines.extend(["\u6682\u65e0\u53ef\u89c6\u5316\u6570\u636e\u3002", ""])
    return lines


def _render_intermediate_markdown(reports: List[Dict[str, str]]) -> List[str]:
    lines = ["## Agent \u4e2d\u95f4\u62a5\u544a", ""]
    if not reports:
        return lines + ["\u6682\u65e0\u4e2d\u95f4\u62a5\u544a\u3002", ""]
    for report in reports:
        lines.extend([f"### {report.get('title', '')}", "", report.get("content", ""), ""])
    return lines


def _render_quality_markdown(quality_report: Dict[str, Any]) -> List[str]:
    lines = ["## \u8d28\u91cf\u8bc4\u5206", ""]
    if not quality_report:
        return lines + ["\u6682\u65e0\u8d28\u91cf\u8bc4\u5206\u3002", ""]
    lines.append(f"- \u603b\u5206: {quality_report.get('overall_score', 'N/A')}")
    dimension_scores = quality_report.get("dimension_scores")
    if isinstance(dimension_scores, dict):
        lines.extend(["", markdown_table([{"dimension": key, "score": value} for key, value in dimension_scores.items()])])
    lines.append("")
    return lines


def markdown_table(rows: List[Dict[str, Any]]) -> str:
    """Render dictionaries as a Markdown table."""
    if not rows:
        return ""
    headers = list(rows[0].keys())
    lines = [
        "| " + " | ".join(_escape_table_cell(header) for header in headers) + " |",
        "| " + " | ".join("---" for _header in headers) + " |",
    ]
    for row in rows:
        lines.append(
            "| "
            + " | ".join(_escape_table_cell(row.get(header, "")) for header in headers)
            + " |"
        )
    return "\n".join(lines)


def _render_docx_bytes(markdown: str) -> bytes:
    try:
        from docx import Document
        from docx.oxml.ns import qn
    except ImportError as exc:
        raise ExportDependencyError(
            "Word export requires python-docx. Install it with: pip install python-docx"
        ) from exc

    document = Document()
    normal_style = document.styles["Normal"]
    normal_style.font.name = "Microsoft YaHei"
    normal_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    in_code = False
    for line in markdown.splitlines():
        stripped = line.strip()
        if stripped == "```json" or stripped == "```":
            in_code = not in_code
            continue
        if not stripped:
            continue
        if in_code:
            document.add_paragraph(stripped)
        elif stripped.startswith("# "):
            document.add_heading(stripped[2:], 0)
        elif stripped.startswith("## "):
            document.add_heading(stripped[3:], 1)
        elif stripped.startswith("### "):
            document.add_heading(stripped[4:], 2)
        elif stripped.startswith("- "):
            document.add_paragraph(stripped[2:], style="List Bullet")
        else:
            document.add_paragraph(stripped)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _render_pdf_bytes(markdown: str, font_path: Optional[str] = None) -> bytes:
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.pdfgen import canvas
    except ImportError as exc:
        raise ExportDependencyError(
            "PDF export requires reportlab. Install it with: pip install reportlab"
        ) from exc

    resolved_font_path = font_path or os.environ.get("EXPORT_CHINESE_FONT_PATH")
    if _contains_cjk(markdown):
        if not resolved_font_path or not Path(resolved_font_path).exists():
            raise ExportDependencyError(
                "Chinese PDF export requires a valid Chinese font path via "
                "EXPORT_CHINESE_FONT_PATH or font_path."
            )
        font_name = "BaziChineseFont"
        pdfmetrics.registerFont(TTFont(font_name, resolved_font_path))
    else:
        font_name = "Helvetica"

    buffer = io.BytesIO()
    page = canvas.Canvas(buffer, pagesize=A4)
    width, height = A4
    margin = 42
    y = height - margin
    line_height = 15
    max_chars = 82
    page.setFont(font_name, 10)

    for raw_line in _markdown_to_plain_lines(markdown):
        for line in _wrap_text(raw_line, max_chars):
            if y < margin:
                page.showPage()
                page.setFont(font_name, 10)
                y = height - margin
            page.drawString(margin, y, line)
            y -= line_height
        if raw_line == "":
            y -= 4

    page.save()
    return buffer.getvalue()


def _markdown_to_plain_lines(markdown: str) -> Iterable[str]:
    in_code = False
    for line in markdown.splitlines():
        stripped = line.strip()
        if stripped == "```json" or stripped == "```":
            in_code = not in_code
            continue
        if in_code:
            yield line
        elif stripped.startswith("#"):
            yield stripped.lstrip("#").strip()
        else:
            yield line.replace("|", "  ")


def _wrap_text(text: str, max_chars: int) -> List[str]:
    if len(text) <= max_chars:
        return [text]
    return [text[index : index + max_chars] for index in range(0, len(text), max_chars)]


def _contains_cjk(text: str) -> bool:
    return any("\u4e00" <= char <= "\u9fff" for char in text)


def _escape_table_cell(value: Any) -> str:
    return str(value).replace("|", "\\|").replace("\n", "<br>")


def _unique_path(path: Path) -> Path:
    if not path.exists():
        return path
    stem = path.stem
    suffix = path.suffix
    for index in range(2, 1000):
        candidate = path.with_name(f"{stem}_{index}{suffix}")
        if not candidate.exists():
            return candidate
    raise FileExistsError(f"Could not create unique path for {path}")
